#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成合成示例 report.docx(与 report.md 同题材,每条内置规则至少埋 2 处问题)。

仅用于生成示例:依赖 python-docx(运行环境的解释器需安装);
qa_check.py 本身不依赖任何第三方库。

源码里的中文串用半角 ,;: 书写,写入 docx 前由 zhs() 统一转为中文语境全角;
「18%,」这处半角逗号是故意保留的 T1 埋点。
"""

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().with_name('report.docx')

_HAN = re.compile('[\u4e00-\u9fff]')
_URL = re.compile(r"(?:https?|ftp)://[A-Za-z0-9._~:/?#\[\]@!$&'()*+,;%=-]+")
_FULL = {',': chr(0xFF0C), ';': chr(0xFF1B), ':': chr(0xFF1A)}
_PLANT = '18%,' + chr(0x5BA2)   # 埋点:半角逗号紧跟「客」


def zhs(s: str) -> str:
    """中文语境的半角 ,;: 转全角;豁免 URL、数字对、盘符,并保护埋点。"""
    s = s.replace(_PLANT, '\x01')
    urls = []

    def stash(m):
        urls.append(m.group())
        return '\x02%d\x02' % (len(urls) - 1)

    t = _URL.sub(stash, s)
    res = []
    for i, ch in enumerate(t):
        if ch in _FULL:
            prev = t[i - 1] if i > 0 else ''
            nxt = t[i + 1] if i + 1 < len(t) else ''
            if prev.isdigit() and nxt.isdigit():
                res.append(ch)
                continue
            if ch == ':' and prev.isascii() and prev.isalpha() and nxt and nxt in '\\/':
                res.append(ch)
                continue
            if _HAN.search(t[max(0, i - 6):i + 7]):
                res.append(_FULL[ch])
                continue
        res.append(ch)
    out = ''.join(res).replace('\x01', _PLANT)
    for k, u in enumerate(urls):                     # 还原被遮蔽的 URL
        out = out.replace('\x02%d\x02' % k, u)
    return out


doc = Document()


def para(text: str) -> None:
    doc.add_paragraph(zhs(text))


def heading(text: str, level: int) -> None:
    doc.add_heading(zhs(text), level=level)


def bold_label(label: str, rest: str) -> None:
    """「加粗标签:」+普通文本(埋 A4 加粗标签列表化问题)。"""
    p = doc.add_paragraph()
    r = p.add_run(zhs(label))
    r.bold = True
    p.add_run(zhs(rest))


def para_deviant_font(text: str, font: str = '楷体') -> None:
    """整段使用与主流不一致的字体(埋 T2 字体不统一问题)。"""
    p = doc.add_paragraph()
    r = p.add_run(zhs(text))
    r.font.name = font
    rfonts = r._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), font)


heading('2026年第三季度交付质量报告', 1)                                     # 第1段
para('编写:张三;复核:李四;归档路径 D:\\company\\reports\\2026q3。')        # 第2段
heading('一、项目概况', 2)                                                   # 第3段
para('本季度团队共完成三个批次的交付,累计处理工单 12,450 件,较去年同期上升 18%,客户满意度 4.6 分。首批次于 2026-07-15 通过验收,第二批次于 2026.8.20 完成上线,第三批次计划 9月30日 收尾。三条业务线中,已有 5条 完成工具化改造。')  # 第4段
para('截至 2026-09-30,系统连续 90 天 无重大故障。磁盘容量 800GB,备份空间 500 GB,内存占用率维持在 62 %左右,整体健康。')  # 第5段
heading('二、质量数据分析', 2)                                               # 第6段
heading('2.1 缺陷分布', 4)          # 埋 T3:H2 后直接 H4                    # 第7段
para_deviant_font('本季度共记录缺陷 320 个,其中严重缺陷 3 个,一般缺陷 45 个,轻微缺陷 272 个,严重缺陷均在 24 小时内闭环。值得注意的是,缺陷集中在接口模块,占六成。')  # 第8段(埋 T2)
para_deviant_font('修复团队由四名工程师与 2名 值班人员组成,平均修复耗时 1.5 天。')  # 第9段(埋 T2)
heading('2.2 近三月说明', 4)                                                 # 第10段
para('数据口径为当月新增,详见 https://reports.company.local/q3/dashboard,如有疑问请联系王经理?')  # 第11段(埋 T1 半角问号)
heading('三、过程改进', 2)                                                   # 第12段
para('首先,我们建立了交付前检查清单,覆盖排版与敏感信息。')                # 第13段(埋 A3)
para('其次,引入了自动化回归,发布 3 次 均无回退。')                        # 第14段(埋 A3)
para('最后,规范了值班交接,详细值班表见内网门户 http://intranet.company.local/duty。')  # 第15段(埋 A3;示例配置的域名)
para('改进效果方面,需求更清晰,分工更明确,交付更顺畅。响应更快,延迟更低,体验更好。变更窗口固定在 22:00。')  # 第16段(埋 A5 两处)
bold_label('进度:', '正常')                                                 # 第17段
bold_label('风险:', '中')                                                   # 第18段
bold_label('结论:', '按期交付')                                             # 第19段(埋 A4 第一组)
heading('四、遗留问题', 2)                                                   # 第20段
heading('4.1 遗留明细', 4)          # 埋 T3:又一处层级跳跃                  # 第21段
para('目前仍有 97200 行 遗留代码未迁移,历史构建产物 57600 个 未清理,GPT模型 需要人工复核生成的段落。值得注意的是,API 网关配额仅剩 40 %。')  # 第22段
para('综上所述,技术债清理要进入下季度计划。内网代号 凤凰计划 的迁移已排期。让我们在下季度——而不是更晚——把 QA检查 前移到构建阶段。')  # 第23段
bold_label('待办一:', '补充灰度方案')                                       # 第24段
bold_label('待办二:', '补齐监控告警')                                       # 第25段
bold_label('待办三:', '更新应急预案')                                       # 第26段(埋 A4 第二组)
heading('五、附录:配置片段', 2)                                             # 第27段
para('接口鉴权(示例,已失效):')                                            # 第28段
para('AWS 访问密钥:AKIA2QKX9LMW7TGZ4PJ2')                                   # 第29段(埋 S1)
para('模型服务密钥:sk-live-9f2K7xQmZ4bN8vLs3RtU')                           # 第30段(埋 S1)
para('代码仓库令牌:ghp_9Xk2mPq4rS7tV1wYz5L8nB6cD0eF3aH')                   # 第31段(埋 S1)
para('鉴权头示例:Bearer eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9')             # 第32段(埋 S1)
para('私钥开头:-----BEGIN RSA PRIVATE KEY-----')                            # 第33段(埋 S1)
para('联系人:王经理,手机 13812345678;备用联系人王经理助理,电话 13998877666,证件号 110101199003077758。')  # 第34段(埋 S2)
para('总而言之,本季度交付质量整体可控——响应及时——风险收敛——团队稳定,相信在大家的共同努力下,交付质量必将迈上新台阶。')  # 第35段(埋 A1/A2/A6)

doc.save(OUT)
print('已生成:' + str(OUT))
